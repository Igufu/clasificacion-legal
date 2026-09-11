"""Stage 1: extract Statement No. / Raw Statement / Attribute / Deontic from PLAN_INSUMO.xlsx.

Reads sheet "DECRETO 31" (row 1 = merged group labels, row 2 = real header,
data starts row 3) and writes outputs/statements.json.

The Raw Statement cell uses **bold** runs to mark the surface phrase that
carries the deontic modality in the text (e.g. "deberán contar", "generará",
"prohíbe la operación"). We capture that as `deontic_surface` so a later
stage can locate it in the official document and tag it. The canonical
deontic word/phrase from the "Deontic" column is stored as `deontic` (may be
bracketed like "[deberá]" when the modal is implied rather than explicit).
"""

import json

import openpyxl

from paths import STATEMENTS_JSON_PATH, XLSX_PATH, ensure_outputs_dir

SHEET_NAME = "DECRETO 31"
HEADER_ROW = 2
REQUIRED_COLUMNS = ["Statement No.", "Raw Statement", "Attribute", "Deontic"]


def find_columns(ws):
    headers = {}
    for cell in ws[HEADER_ROW]:
        if cell.value is None:
            continue
        name = str(cell.value).strip()
        if name:
            headers[name] = cell.column

    missing = [name for name in REQUIRED_COLUMNS if name not in headers]
    if missing:
        raise ValueError(f"Missing expected column(s) in header row {HEADER_ROW}: {missing}")

    return {name: headers[name] for name in REQUIRED_COLUMNS}


def is_blank(value):
    if value is None:
        return True
    if isinstance(value, str):
        return value.strip() == ""
    text = str(value)
    return text.strip() == ""


def extract_bold_surface(cell_value):
    """Concatenate bold-formatted runs inside a Raw Statement cell.

    With `rich_text=True`, an unformatted cell is still a plain str (no runs
    to inspect) and returns "". A formatted cell is a CellRichText whose
    runs each carry an InlineFont; we keep only the runs with `font.b` set.
    """
    if cell_value is None or isinstance(cell_value, str):
        return ""
    parts = []
    for run in cell_value:
        text = getattr(run, "text", None)
        if text is None:
            continue
        font = getattr(run, "font", None)
        if font is not None and font.b:
            parts.append(text)
    return "".join(parts)


def extract_statements(ws, columns):
    statements = []
    for row in range(HEADER_ROW + 1, ws.max_row + 1):
        raw_cell = ws.cell(row=row, column=columns["Raw Statement"]).value
        if is_blank(raw_cell):
            continue

        stmt_no_value = ws.cell(row=row, column=columns["Statement No."]).value
        if is_blank(stmt_no_value):
            continue
        stmt_no = int(stmt_no_value)

        attribute_value = ws.cell(row=row, column=columns["Attribute"]).value
        attribute = None if is_blank(attribute_value) else str(attribute_value)

        deontic_value = ws.cell(row=row, column=columns["Deontic"]).value
        deontic = None if is_blank(deontic_value) else str(deontic_value)

        deontic_surface_raw = extract_bold_surface(raw_cell)
        deontic_surface = deontic_surface_raw if deontic_surface_raw.strip() else None

        statements.append(
            {
                "stmt_no": stmt_no,
                "raw_statement": str(raw_cell),
                "attribute": attribute,
                "deontic": deontic,
                "deontic_surface": deontic_surface,
            }
        )

    return statements


def main():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True, rich_text=True)
    ws = wb[SHEET_NAME]

    columns = find_columns(ws)
    statements = extract_statements(ws, columns)

    ensure_outputs_dir()
    with open(STATEMENTS_JSON_PATH, "w", encoding="utf-8") as f:
        json.dump(statements, f, ensure_ascii=False, indent=2)

    print(f"Extracted {len(statements)} statements -> {STATEMENTS_JSON_PATH}")


if __name__ == "__main__":
    main()
