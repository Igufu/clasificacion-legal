"""Stage 1: extract Statement No. / Raw Statement / Attribute from PLAN_INSUMO.xlsx.

Reads sheet "DECRETO 31" (row 1 = merged group labels, row 2 = real header,
data starts row 3) and writes statements.json.
"""

import json
from pathlib import Path

import openpyxl

BASE_DIR = Path(__file__).resolve().parent
XLSX_PATH = BASE_DIR / "PLAN_INSUMO.xlsx"
SHEET_NAME = "DECRETO 31"
HEADER_ROW = 2
JSON_PATH = BASE_DIR / "statements.json"

REQUIRED_COLUMNS = ["Statement No.", "Raw Statement", "Attribute"]


def find_columns(ws):
    headers = {}
    for cell in ws[HEADER_ROW]:
        if cell.value is None:
            continue
        name = str(cell.value).strip()
        if name:
            headers[name] = cell.column

    missing = [name for name in REQUIRED_COLUMNS if name not in headers]
    if missing:
        raise ValueError(f"Missing expected column(s) in header row {HEADER_ROW}: {missing}")

    return {name: headers[name] for name in REQUIRED_COLUMNS}


def is_blank(value):
    return value is None or (isinstance(value, str) and value.strip() == "")


def extract_statements(ws, columns):
    statements = []
    for row in range(HEADER_ROW + 1, ws.max_row + 1):
        raw_statement = ws.cell(row=row, column=columns["Raw Statement"]).value
        if is_blank(raw_statement):
            continue

        stmt_no_value = ws.cell(row=row, column=columns["Statement No."]).value
        if is_blank(stmt_no_value):
            continue
        stmt_no = int(stmt_no_value)

        attribute_value = ws.cell(row=row, column=columns["Attribute"]).value
        attribute = None if is_blank(attribute_value) else str(attribute_value)

        statements.append(
            {
                "stmt_no": stmt_no,
                "raw_statement": str(raw_statement),
                "attribute": attribute,
            }
        )

    return statements


def main():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb[SHEET_NAME]

    columns = find_columns(ws)
    statements = extract_statements(ws, columns)

    with open(JSON_PATH, "w", encoding="utf-8") as f:
        json.dump(statements, f, ensure_ascii=False, indent=2)

    print(f"Extracted {len(statements)} statements -> {JSON_PATH}")


if __name__ == "__main__":
    main()
