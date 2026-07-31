import json
from pathlib import Path
import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Font, PatternFill

OUTPUT_DIR = Path(__file__).resolve().parent.parent / "output"
INPUT_FILE = OUTPUT_DIR / "extraction_results.json"
OUTPUT_FILE = OUTPUT_DIR / "extraction_results.xlsx"
CATEGORIES_PATH = Path(__file__).resolve().parent / "data" / "categories.csv"
SHEET_NAME = "Statements"

# Per-column widths (chars). Long free-text fields get more room. Anything
# without an explicit entry (including any never-anticipated extra key found
# in the data) falls back to DEFAULT_WIDTH.
COLUMN_WIDTHS = {
    "Section Marker": 32,
    "Statement No.": 8,
    "Raw Statement": 70,
    "Statement Type": 14,
    "Attribute": 28,
    "Attribute Property": 40,
    "Deontic": 12,
    "Aim": 16,
    "Direct Object (Content)": 32,
    "Direct Object (Reference to statement)": 16,
    "Direct Object Property": 40,
    "Indirect Object (Content)": 32,
    "Indirect Object (Reference to statement)": 16,
    "Indirect Object Property": 40,
    "Activation Condition (Content)": 50,
    "Activation Condition (Reference to statement)": 16,
    "Execution Constraint (Content)": 40,
    "Execution Constraint (Reference to statement)": 16,
    "Or else": 40,
    "Or else (Forward reference)": 16,
    "Logical combination": 18,
    "Invoking statement(s)": 14,
    "Constituted Entity": 32,
    "Constituted Entity Property": 40,
    "Constituting Function": 24,
    "Constituted Properties": 50,
    "Constituted Properties Property": 40,
}
DEFAULT_WIDTH = 25


def load_records(path: Path) -> list[dict]:
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def load_canonical_fields(path: Path) -> list[str]:
    """Column order sourced straight from categories.csv — the same file
    prompt_builder.build_response_schema() reads to build the model's output
    schema — so the two can never drift apart the way the old hand-maintained
    COLUMN_ORDER list did. Row order is preserved; a name repeated across the
    regulative/constitutive sections of the CSV (e.g. "Deontic") collapses to
    its first occurrence, same dedup rule build_response_schema uses."""
    df = pd.read_csv(path)
    df.columns = ["Coding Fields", "Description"]
    fields = []
    for name in df["Coding Fields"]:
        name = str(name).strip()
        if name and name not in fields:
            fields.append(name)
    return fields


def order_columns(records: list[dict], canonical_fields: list[str]) -> list[str]:
    """Every canonical Coding Field gets a column even if no record in this
    particular run happened to populate it, so the sheet always shows the
    full coding framework. Truly unexpected keys (not in categories.csv) are
    appended at the end so nothing is silently dropped."""
    seen = {k for r in records for k in r.keys()}
    extras = sorted(seen - set(canonical_fields))
    return canonical_fields + extras


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    records = load_records(INPUT_FILE)
    canonical_fields = load_canonical_fields(CATEGORIES_PATH)
    columns = order_columns(records, canonical_fields)
    df = pd.DataFrame(records).reindex(columns=columns)

    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=SHEET_NAME)
        ws = writer.sheets[SHEET_NAME]

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        body_align = Alignment(vertical="top", wrap_text=True)

        for col_idx, name in enumerate(df.columns, start=1):
            letter = get_column_letter(col_idx)
            ws.column_dimensions[letter].width = COLUMN_WIDTHS.get(name, DEFAULT_WIDTH)
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
            for cell in row:
                cell.alignment = body_align

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    print(f"Wrote {len(df)} rows to {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
